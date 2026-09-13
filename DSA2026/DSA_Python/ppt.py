from pptx import Presentation

prs = Presentation("KisanSetuAI.pptx")

for slide in prs.slides:
    for shape in list(slide.shapes):
        # Text frame me gamma dhoond kar delete karna
        if shape.has_text_frame and "gamma" in shape.text.lower():
            sp = shape._element
            sp.getparent().remove(sp)

prs.save("KisanSetuAI_Clean.pptx")
print("Badge successfully removed from PPTX!")